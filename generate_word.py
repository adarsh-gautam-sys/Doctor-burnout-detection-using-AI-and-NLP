import sys
import re
import os
from docx import Document
from docx.shared import Inches

def create_docx_from_md(md_file, docx_file):
    doc = Document()
    
    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('# '):
            doc.add_heading(line[2:], level=0)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('!['):
            # Parse image
            m = re.search(r'!\[.*?\]\((.*?)\)', line)
            if m:
                img_path = m.group(1)
                # Fix file:/// protocol
                if img_path.startswith('file:///'):
                    img_path = img_path[8:]
                # Fix slashes for windows
                img_path = img_path.replace('/', '\\')
                img_path = img_path.replace('%20', ' ')
                
                if os.path.exists(img_path):
                    try:
                        doc.add_picture(img_path, width=Inches(6.0))
                    except Exception as e:
                        doc.add_paragraph(f"[Image could not be embedded (format not supported by docx): {img_path}]")
                else:
                    doc.add_paragraph(f"[Image not found at path: {img_path}]")
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            text = line[2:]
            _add_formatted_text(p, text)
        elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. ') or line.startswith('4. ') or line.startswith('5. ') or line.startswith('6. '):
            p = doc.add_paragraph(style='List Number')
            text = line[3:]
            _add_formatted_text(p, text)
        elif line.startswith('***') or line.startswith('---'):
            pass # skip horizontal lines
        else:
            p = doc.add_paragraph()
            _add_formatted_text(p, line)
            
    doc.save(docx_file)

def _add_formatted_text(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            subparts = re.split(r'(\*.*?\*)', part)
            for sp in subparts:
                if sp.startswith('*') and sp.endswith('*'):
                    run = paragraph.add_run(sp[1:-1])
                    run.italic = True
                else:
                    bparts = re.split(r'(`.*?`)', sp)
                    for bp in bparts:
                        if bp.startswith('`') and bp.endswith('`'):
                            run = paragraph.add_run(bp[1:-1])
                            # Optionally add code font style
                        else:
                            paragraph.add_run(bp)

if __name__ == '__main__':
    create_docx_from_md(sys.argv[1], sys.argv[2])
